#!/usr/bin/env python3
"""
Extract raw text from an existing resume PDF or DOCX. Deliberately does NOT
try to structure the output into sections/bullets -- that step needs
judgment (a resume's visual layout doesn't map 1:1 onto its logical
structure, e.g. a two-column layout or a skills table), and is handled by
Claude reading this script's output, not by more parsing code. This script's
only job is reliable text extraction -- the part that's actually mechanical.

Usage:
  python3 extract_text.py <resume.pdf|resume.docx>
"""
import subprocess
import sys
from pathlib import Path


def extract_pdf(path):
    result = subprocess.run(
        ["pdftotext", "-layout", str(path), "-"],
        capture_output=True, text=True, check=True,
    )
    return result.stdout


def extract_docx(path):
    import docx  # python-docx
    doc = docx.Document(str(path))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    path = Path(sys.argv[1])
    if not path.exists():
        print(f"File not found: {path}", file=sys.stderr)
        sys.exit(1)

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = extract_pdf(path)
    elif suffix == ".docx":
        text = extract_docx(path)
    else:
        print(f"Unsupported file type: {suffix} (only .pdf and .docx)", file=sys.stderr)
        sys.exit(1)

    print(text)


if __name__ == "__main__":
    main()
